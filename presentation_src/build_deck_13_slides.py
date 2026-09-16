"""Builds the GymVerse evaluation deck (presentation.pptx) with python-pptx.

Every number comes from measured results:
  bench-results.json  index ablation benchmark (synthetic 50k members / 1M check-ins)
  e2e-results.json    live end-to-end run against the real database (116 checks)
Set REPO_URL (and optionally CONTACT_EMAIL) to replace the placeholders with real links and
QR codes once the code is pushed.
"""
import io
import json
import os
import sys
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "presentation.pptx"
REPO_URL = os.environ.get("REPO_URL", "").strip()
CONTACT_EMAIL = os.environ.get("CONTACT_EMAIL", "").strip()

BENCH = json.loads((HERE / "bench-results.json").read_text())
E2E = json.loads((HERE / "e2e-results.json").read_text())

# ------------------------------------------------------------------ theme
DARK = RGBColor(0x11, 0x18, 0x27)
BODY = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LINE = RGBColor(0xD1, 0xD5, 0xDB)
SOFT = RGBColor(0xF3, 0xF4, 0xF6)
MID = RGBColor(0x9C, 0xA3, 0xAF)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT_TINT = RGBColor(0xDB, 0xEA, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
MONO = "Consolas"
EMOJI = "Segoe UI Emoji"
TOTAL_SLIDES = 13

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, lines, size=18, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.0, space_after=0):
    """Text box. `lines` is a string or a list of strings / (string, overrides) tuples."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    items = [lines] if isinstance(lines, str) else lines
    for i, item in enumerate(items):
        content, opts = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = spacing
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = content
        f = run.font
        f.name = opts.get("font", font)
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", False)
        f.color.rgb = opts.get("color", color)
    return box


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.08, line_width=1.25, dash=False,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_width)
        if dash:
            s.line.dash_style = 4  # dash
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, content, fill, line, color, size=18, bold=True, radius=0.1):
    s = box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = content
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def circle_number(slide, x, y, d, content, fill=ACCENT, color=WHITE, size=22):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = content
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
    return s


def new_slide(number, title, tag=None, notes=""):
    slide = prs.slides.add_slide(BLANK)
    if title:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.38), Inches(0.55), Inches(0.08), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        bar.shadow.inherit = False
        text(slide, 0.6, 0.42, 10.3, 1.25, title, size=34, color=DARK, bold=True, anchor=MSO_ANCHOR.TOP, spacing=0.95)
    if tag:
        label_box(slide, 11.15, 0.5, 1.65, 0.46, tag, fill=WHITE, line=ACCENT, color=ACCENT, size=14, radius=0.5)
    if number:
        text(slide, 12.2, 6.95, 0.7, 0.4, str(number), size=14, color=MUTED, align=PP_ALIGN.RIGHT)
        text(slide, 0.6, 6.95, 4.0, 0.4, "GymVerse · DBMS Mini Project", size=14, color=MUTED)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def qr_picture(slide, url, x, y, size):
    img = qrcode.QRCode(border=1, box_size=10)
    img.add_data(url)
    img.make(fit=True)
    buf = io.BytesIO()
    img.make_image(fill_color="#111827", back_color="white").save(buf, format="PNG")
    buf.seek(0)
    return slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(size), Inches(size))


def qr_or_placeholder(slide, x, y, size, placeholder):
    if REPO_URL:
        qr_picture(slide, REPO_URL, x, y, size)
    else:
        s = box(slide, x, y, size, size, fill=SOFT, line=MID, dash=True, radius=0.04)
        tf = s.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = placeholder
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(18), True, DARK


def set_chart_fonts(chart, size=18):
    from pptx.enum.chart import XL_AXIS_CROSSES
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = BODY
    # No automatic title (a single-series chart gets its series name as a title otherwise),
    # and the category axis sits at the value axis minimum, so bars below 1 on a log scale
    # still grow upwards from the baseline.
    chart.has_title = False
    chart.value_axis.crosses = XL_AXIS_CROSSES.MINIMUM


# ------------------------------------------------------------------ data
bench_q = {q["key"]: q for q in BENCH["queries"]}
bench_order = [("payments_member", "Payment history"), ("subscriptions_list", "Subscriptions list"),
               ("attendance_log", "Attendance log"), ("member_search", "Member search")]
speedups = {k: bench_q[k]["results"]["before16"]["median_ms"] / bench_q[k]["results"]["after16"]["median_ms"]
            for k, _ in bench_order}
best = max(speedups.values())
vol = BENCH["volumes"]

groups = [
    ("Accounts & security", ["Login", "Register member", "Trainer approval", "Account security", "Access control"]),
    ("Training & classes", ["Exercises", "Workouts", "Classes"]),
    ("Dashboards, health & UI", ["Dashboards & reports", "Health", "Frontend"]),
    ("Members & plans", ["Members", "Plans"]),
    ("Payments & subscriptions", ["Subscriptions & payments"]),
    ("Attendance & rosters", ["Attendance", "Trainer rosters"]),
]
area_pass = {}
for r in E2E["results"]:
    area_pass.setdefault(r["area"], [0, 0])
    area_pass[r["area"]][0 if r["ok"] else 1] += 1
group_counts = [(name, sum(area_pass[a][0] for a in areas), sum(sum(area_pass[a]) for a in areas)) for name, areas in groups]
e2e_passed = sum(1 for r in E2E["results"] if r["ok"])
e2e_total = len(E2E["results"])
assert sum(t for _, _, t in group_counts) == e2e_total, "every end-to-end check must be in a group"

TEAM = [("NYASHADAM SRIKAR", "2510030006"), ("MALLALA SRI CHARAN", "2510030246"),
        ("YOGENDRA UNGARALA", "2510030147"), ("GORANTAL PRANAV", "2510030008")]

# ================================================================== 1 Title
s = new_slide(None, None, notes=(
    "Introduce GymVerse: a PostgreSQL-backed gym management system with a web app and an AI assistant. "
    "Team of four, Database Management Systems mini project."))
text(s, 0.8, 0.95, 7.2, 0.5, "DATABASE MANAGEMENT SYSTEMS · MINI PROJECT", size=18, color=ACCENT, bold=True)
text(s, 0.8, 1.45, 7.2, 1.3, "GymVerse", size=66, color=DARK, bold=True)
text(s, 0.8, 2.8, 7.0, 1.1, "The Digital Transformation of Fitness Management", size=28, color=DARK)
text(s, 0.8, 4.1, 6.6, 1.0, "One connected, self-protecting database for gym operations", size=20, color=MUTED)
box(s, 8.35, 0.95, 4.4, 4.95, fill=SOFT, line=None, radius=0.05)
text(s, 8.7, 1.2, 3.8, 0.5, "Team", size=20, color=DARK, bold=True)
for i, (name, roll) in enumerate(TEAM):
    y = 1.85 + i * 1.0
    text(s, 8.7, y, 3.9, 0.45, name, size=18, color=DARK, bold=True)
    text(s, 8.7, y + 0.4, 3.9, 0.45, roll, size=18, color=MUTED)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.95), Inches(0.03))
line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background(); line.shadow.inherit = False
text(s, 0.8, 6.35, 11.9, 0.5, "Department of CSE  ·  KLH  ·  Faculty Guide: P LALITHA", size=18, color=MUTED)

# ================================================================== 2 Problem
s = new_slide(2, "Scattered records cause duplicates and double bookings", notes=(
    "Many gyms keep separate registers and spreadsheets. Nothing stops the same member being entered twice, "
    "an expired plan going unnoticed, or two people taking the last class seat at the same time. "
    "Spreadsheets store data but cannot enforce these rules. GymVerse moves the rules into the database."))
box(s, 0.6, 1.95, 5.6, 4.5, fill=SOFT, line=None, radius=0.05)
text(s, 0.95, 2.2, 5.0, 0.6, "Registers & spreadsheets", size=24, color=BODY, bold=True)
box(s, 7.15, 1.95, 5.6, 4.5, fill=ACCENT_TINT, line=ACCENT, radius=0.05, line_width=1.5)
text(s, 7.5, 2.2, 5.0, 0.6, "GymVerse database", size=24, color=ACCENT, bold=True)
problems = ["Same member stored twice", "Expired plans go unnoticed", "Two people get the last class seat"]
fixes = ["One record per member", "Expiry tracked automatically", "Capacity enforced on every booking"]
for i, (p_, f_) in enumerate(zip(problems, fixes)):
    y = 3.1 + i * 1.0
    text(s, 0.95, y, 0.5, 0.6, "✕", size=24, color=MID, bold=True, font="Segoe UI Symbol")
    text(s, 1.5, y + 0.02, 4.5, 0.8, p_, size=20, color=BODY)
    text(s, 7.5, y, 0.5, 0.6, "✓", size=24, color=ACCENT, bold=True, font="Segoe UI Symbol")
    text(s, 8.05, y + 0.02, 4.5, 0.8, f_, size=20, color=DARK)
arrow(s, 6.3, 4.2, 7.05, 4.2, color=ACCENT, width=3)

# ================================================================== 3 Novelty: comparison
s = new_slide(3, "Only GymVerse enforces gym rules inside the database", tag="Novelty", notes=(
    "Rows are approaches; columns are capabilities. 'Partial' for spreadsheets: data validation, sheet-level sharing "
    "and version history exist but do not enforce gym rules or per-record access. "
    "[ADD APPROACH]: add another existing system you compared against, and fill in its row."))
cols = ["Approach", "One source of truth", "Rules enforced in database", "Safe with concurrent users",
        "Per-record access", "Audit trail", "AI on live data"]
rows = [
    ("Paper registers", ["✕", "✕", "✕", "✕", "✕", "✕"]),
    ("Spreadsheets", ["✕", "Partial", "✕", "Partial", "Partial", "✕"]),
    ("[ADD APPROACH]", ["[ADD]"] * 6),
    ("GymVerse (ours)", ["✓"] * 6),
]
gf = s.shapes.add_table(len(rows) + 1, len(cols), Inches(0.6), Inches(1.9), Inches(12.1), Inches(4.6))
tbl = gf.table
tbl.first_row = False
tbl.horz_banding = False
tbl.columns[0].width = Inches(2.95)
for c in range(1, len(cols)):
    tbl.columns[c].width = Inches((12.1 - 2.95) / 6)
tbl.rows[0].height = Inches(1.4)
for r in range(1, len(rows) + 1):
    tbl.rows[r].height = Inches(0.8)


def cell(r, c, content, size=18, bold=False, color=BODY, fill=WHITE, align=PP_ALIGN.CENTER, font=FONT):
    ce = tbl.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    ce.margin_left = ce.margin_right = Inches(0.08)
    tf = ce.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.text = ""
    run = p.add_run()
    run.text = content
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = font, Pt(size), bold, color


for c, h in enumerate(cols):
    cell(0, c, h, size=18, bold=True, color=DARK, fill=SOFT, align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
for r, (name, marks) in enumerate(rows, start=1):
    ours = name.startswith("GymVerse")
    fill = ACCENT_TINT if ours else WHITE
    cell(r, 0, name, size=20, bold=True, color=ACCENT if ours else DARK, fill=fill, align=PP_ALIGN.LEFT)
    for c, m in enumerate(marks, start=1):
        symbol = m in ("✓", "✕")
        cell(r, c, m, size=26 if symbol else 18, bold=ours or symbol,
             color=ACCENT if ours else (MID if m == "✕" else BODY), fill=fill,
             font="Segoe UI Symbol" if symbol else FONT)

# ================================================================== 4 Novelty: contributions
s = new_slide(4, "Three contributions set GymVerse apart", tag="Novelty", notes=(
    "1) Business rules are enforced by PostgreSQL itself: partial unique indexes (one active subscription, one open check-in), "
    "row locks for bookings and subscriptions, triggers for capacity and auditing. "
    "2) Four roles, checked per record: a member reaches only their own data, a trainer only their roster; sessions can be revoked. "
    "3) The AI assistant reads the live plan catalogue and falls back across models when one is out of quota."))
contrib = [
    ("Rules live in the database", "Unique indexes, row locks and triggers stop duplicates and overbooking."),
    ("Every record checks who is asking", "Four roles: members see only themselves, trainers only their roster."),
    ("AI grounded on live data", "Answers from the real plan catalogue, with automatic model fallback."),
]
for i, (head, desc) in enumerate(contrib):
    x = 0.6 + i * 4.15
    box(s, x, 1.95, 3.85, 4.55, fill=WHITE, line=LINE, radius=0.05)
    circle_number(s, x + 0.35, 2.3, 0.75, str(i + 1))
    text(s, x + 0.35, 3.3, 3.2, 1.1, head, size=22, color=DARK, bold=True, spacing=0.95)
    text(s, x + 0.35, 4.45, 3.2, 1.8, desc, size=18, color=BODY)

# ================================================================== 5 Novelty: architecture
s = new_slide(5, "Every request passes four guards before touching data", tag="Novelty", notes=(
    "Browser requests go through rate limiting, session checks (JWT plus server-side revocation), role and per-record access, "
    "and input validation. Controllers run multi-row writes in transactions with audit entries. PostgreSQL enforces the rules "
    "with indexes, row locks and triggers; schema changes ship as versioned migrations. The AI service tries Gemini models "
    "in order, then Groq; chat history in MongoDB is optional. Blue = novel in GymVerse, grey = standard."))
label_box(s, 0.5, 2.05, 1.8, 1.0, "React app", fill=SOFT, line=LINE, color=BODY)
arrow(s, 2.32, 2.55, 2.68, 2.55)
guards = [("Rate limit", False), ("Session check", True), ("Role + record access", True), ("Input validation", False)]
gx = [2.7, 4.55, 6.4, 8.25]
for i, ((name, novel), x) in enumerate(zip(guards, gx)):
    label_box(s, x, 2.05, 1.6, 1.0, name,
              fill=ACCENT if novel else SOFT, line=ACCENT if novel else LINE,
              color=WHITE if novel else BODY, size=18)
    if i < 3:
        arrow(s, x + 1.62, 2.55, gx[i + 1] - 0.02, 2.55)
text(s, 2.7, 3.12, 3.6, 0.45, "Express API: four guards", size=18, color=MUTED)
arrow(s, 9.05, 3.07, 9.05, 3.73)
label_box(s, 6.4, 3.75, 3.45, 0.85, "Controllers + transactions", fill=SOFT, line=LINE, color=BODY)
box(s, 10.25, 1.75, 2.6, 4.65, fill=WHITE, line=DARK, radius=0.05, line_width=1.5)
text(s, 10.25, 1.82, 2.6, 0.5, "PostgreSQL", size=20, color=DARK, bold=True, align=PP_ALIGN.CENTER)
chips = [("Unique & partial indexes", True), ("Row locks", True), ("Audit + capacity triggers", True), ("Versioned migrations", False)]
for i, (name, novel) in enumerate(chips):
    label_box(s, 10.4, 2.4 + i * 0.97, 2.3, 0.8, name,
              fill=ACCENT if novel else SOFT, line=ACCENT if novel else LINE,
              color=WHITE if novel else BODY, size=18)
arrow(s, 9.87, 4.17, 10.23, 4.17)
label_box(s, 6.4, 5.25, 1.65, 0.9, "AI model fallback", fill=ACCENT, line=ACCENT, color=WHITE, size=18)
label_box(s, 8.2, 5.25, 1.65, 0.9, "Chat history (MongoDB)", fill=SOFT, line=LINE, color=BODY, size=18)
arrow(s, 7.22, 4.62, 7.22, 5.23)
arrow(s, 9.02, 4.62, 9.02, 5.23)
label_box(s, 4.2, 5.25, 1.95, 0.9, "Gemini → Groq", fill=SOFT, line=LINE, color=BODY, size=18)
arrow(s, 6.38, 5.7, 6.17, 5.7)
legend_y = 6.45
box(s, 0.6, legend_y + 0.08, 0.3, 0.3, fill=ACCENT, line=None, radius=0.2)
text(s, 1.0, legend_y, 2.8, 0.45, "Novel in GymVerse", size=18, color=BODY)
box(s, 3.8, legend_y + 0.08, 0.3, 0.3, fill=SOFT, line=LINE, radius=0.2)
text(s, 4.2, legend_y, 2.8, 0.45, "Standard component", size=18, color=BODY)

# ================================================================== 6 Adaptability
s = new_slide(6, "The same design adapts from one PC to gym chains", tag="Adaptability", notes=(
    "The schema models generic concepts: people, plans, sessions with capacity, bookings, payments. "
    "Studios and college sports centres fit as-is; clinics map sessions to appointments; chains need a branch column "
    "(not built yet). The benchmark shows the same queries stay under 10 ms at 50,000 members and 1,000,000 check-ins. "
    "Plans, AI models, email, chat history and production mode are configuration; integrity rules, record access and "
    "the audit trail are the fixed core."))
cards = [
    ("🏢", "Gym chains", "Needs a branch column"),
    ("🧘", "Studios", "Yoga & dance: works as-is"),
    ("🩺", "Physio clinics", "Sessions map to appointments"),
    ("🎓", "College gyms", "Students as members"),
    ("💻", "Low-resource gyms", "One PC; AI and email optional"),
    ("📈", "Large data volumes", f"Under 10 ms at {vol['attendance'] // 1_000_000}M check-ins"),
]
for i, (icon, head, cap) in enumerate(cards):
    x = 0.6 + (i % 3) * 4.15
    y = 1.9 + (i // 3) * 1.72
    box(s, x, y, 3.85, 1.52, fill=WHITE, line=LINE, radius=0.08)
    text(s, x + 0.15, y + 0.3, 0.85, 0.9, icon, size=30, font=EMOJI, align=PP_ALIGN.CENTER)
    text(s, x + 1.05, y + 0.18, 2.7, 0.5, head, size=20, color=DARK, bold=True)
    text(s, x + 1.05, y + 0.7, 2.7, 0.75, cap, size=18, color=MUTED)
box(s, 0.6, 5.4, 5.95, 1.35, fill=ACCENT_TINT, line=ACCENT, radius=0.08)
text(s, 0.85, 5.5, 5.5, 0.5, "Configurable, no code change", size=20, color=ACCENT, bold=True)
text(s, 0.85, 6.0, 5.5, 0.7, "Plans · AI models · email · chat history · production mode", size=18, color=BODY)
box(s, 6.75, 5.4, 5.95, 1.35, fill=SOFT, line=None, radius=0.08)
text(s, 7.0, 5.5, 5.5, 0.5, "Fixed core", size=20, color=DARK, bold=True)
text(s, 7.0, 6.0, 5.5, 0.7, "Integrity rules · record access · audit trail", size=18, color=BODY)

# ================================================================== 7 Evaluation: setup
member_scale = vol["members"] // 50
s = new_slide(7, f"Tested four ways, on {member_scale:,}× more members than the seed", tag="Evaluation", notes=(
    f"Seed data: 50 members, 300 visits. Benchmark: a throwaway PostgreSQL {BENCH.get('postgres', '')} database built from the real "
    f"migrations with synthetic volumes ({vol['members']:,} members, {vol['subscriptions']:,} subscriptions, {vol['payments']:,} payments, "
    f"{vol['attendance']:,} check-ins) on a single Windows PC. Each of the app's own list and search queries ran {BENCH['runs']} times "
    "per index setup with EXPLAIN ANALYZE; we report the median. Functional checks: Jest unit and route tests, verify_stack.js on a "
    "fresh database, and a live end-to-end run against the real database."))
setup = [
    ("Data", ["Seed: 50 members, 300 visits", f"Benchmark: {vol['members'] // 1000}k members, {vol['attendance'] // 1_000_000}M check-ins (synthetic)"]),
    ("Baselines", ["No supporting indexes", "Indexes before migration 16"]),
    ("Metrics", [f"Median of {BENCH['runs']} query runs", "Pass rate of every check"]),
]
for i, (head, items) in enumerate(setup):
    x = 0.6 + i * 4.15
    box(s, x, 1.95, 3.85, 3.55, fill=WHITE, line=LINE, radius=0.05)
    text(s, x + 0.3, 2.15, 3.3, 0.55, head, size=24, color=ACCENT, bold=True)
    text(s, x + 0.3, 2.85, 3.3, 2.5, [(t, {"space_after": 14}) for t in items], size=20, color=BODY)
layers = ["Unit & route tests", "Fresh-DB stack check", "Live end-to-end", "Index benchmark"]
for i, name in enumerate(layers):
    label_box(s, 0.6 + i * 3.1, 5.8, 2.85, 0.75, name, fill=SOFT, line=None, color=DARK, size=18)

# ================================================================== 8 Evaluation: benchmark (ablation)
s = new_slide(8, f"New indexes make key queries up to {best:.0f}× faster", tag="Evaluation", notes=(
    "Ablation: the same four queries under three index setups, median of 15 runs, log scale. "
    + " ".join(f"{label}: {bench_q[k]['results']['none']['median_ms']:.1f} / {bench_q[k]['results']['before16']['median_ms']:.1f} / "
               f"{bench_q[k]['results']['after16']['median_ms']:.2f} ms." for k, label in bench_order)
    + " Before migration 16 the attendance log was slower than with no indexes; both used a full scan, so that gap is timing noise. "
    "Synthetic data on one PC: absolute times will differ on other hardware."))
cd = CategoryChartData()
cd.categories = [label for _, label in bench_order]
cd.add_series("No indexes", [bench_q[k]["results"]["none"]["median_ms"] for k, _ in bench_order])
cd.add_series("Before migration 16", [bench_q[k]["results"]["before16"]["median_ms"] for k, _ in bench_order])
cd.add_series("GymVerse (after 16)", [bench_q[k]["results"]["after16"]["median_ms"] for k, _ in bench_order])
gframe = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.8), Inches(8.9), Inches(5.05), cd)
chart = gframe.chart
set_chart_fonts(chart, 18)
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.include_in_layout = False
chart.legend.font.size = Pt(18)
plot = chart.plots[0]
plot.gap_width = 60
plot.overlap = -10
for ser, color in zip(plot.series, [LINE, MID, ACCENT]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = color
    ser.format.line.fill.background()
ours = plot.series[2]
ours.data_labels.show_value = True
ours.data_labels.number_format = '0.0" ms"'
ours.data_labels.number_format_is_linked = False
ours.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ours.data_labels.font.size = Pt(16)
ours.data_labels.font.bold = True
ours.data_labels.font.color.rgb = ACCENT
va = chart.value_axis
scaling = va._element.find(qn("c:scaling"))
scaling.insert(0, scaling.makeelement(qn("c:logBase"), {"val": "10"}))
va.minimum_scale = 0.1
va.maximum_scale = 1000
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = SOFT
va.format.line.fill.background()
va.tick_labels.font.size = Pt(16)
va.tick_labels.font.color.rgb = MUTED
va.has_title = True
va.axis_title.text_frame.text = "Median time (ms, log scale)"
va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
va.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
chart.category_axis.tick_labels.font.size = Pt(18)
chart.category_axis.format.line.color.rgb = LINE
text(s, 9.55, 1.95, 3.5, 0.5, "Speed-up after 16", size=18, color=MUTED, bold=True)
for i, (k, label) in enumerate(sorted(bench_order, key=lambda kv: -speedups[kv[0]])):
    y = 2.5 + i * 0.9
    text(s, 9.55, y, 1.45, 0.7, f"{speedups[k]:.0f}×" if speedups[k] >= 10 else f"{speedups[k]:.1f}×", size=32, color=ACCENT, bold=True)
    text(s, 10.95, y + 0.12, 2.2, 0.6, label, size=18, color=BODY)
text(s, 9.55, 6.15, 3.6, 0.5, f"{vol['members'] // 1000}k members, {vol['attendance'] // 1_000_000}M check-ins", size=18, color=MUTED)

# ================================================================== 9 Evaluation: functional results
s = new_slide(9, f"All {e2e_passed} end-to-end checks pass across every role", tag="Evaluation", notes=(
    f"Live run against the real database: {e2e_passed} of {e2e_total} checks, each write confirmed with a direct database query. "
    "Unit and route tests: 235 passed, 1 skipped. verify_stack.js on a fresh database: 71 of 71. Frontend tests: 7 of 7. "
    "No 5xx responses during the run. An earlier run found a real bug (a migration depended on a function older databases lacked); "
    "migration 17 fixed it before these numbers were taken."))
cd = CategoryChartData()
cd.categories = [g for g, _, _ in group_counts]
cd.add_series("Checks passed", [p for _, p, _ in group_counts])
gframe = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.85), Inches(7.9), Inches(4.95), cd)
chart = gframe.chart
set_chart_fonts(chart, 18)
chart.has_legend = False
plot = chart.plots[0]
plot.gap_width = 55
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = ACCENT
ser.format.line.fill.background()
ser.data_labels.show_value = True
ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser.data_labels.font.size = Pt(18)
ser.data_labels.font.bold = True
ser.data_labels.font.color.rgb = DARK
ca = chart.category_axis
ca.reverse_order = True
ca.tick_labels.font.size = Pt(18)
ca.format.line.color.rgb = LINE
va = chart.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = max(t for _, _, t in group_counts) * 1.2
va.minimum_scale = 0
tiles = [("235", "unit & route tests pass"), ("71/71", "fresh-DB stack checks"),
         ("7/7", "frontend tests"), ("0", "server errors (5xx)")]
for i, (big, cap) in enumerate(tiles):
    x = 8.75 + (i % 2) * 2.1
    y = 1.95 + (i // 2) * 2.45
    box(s, x, y, 1.95, 2.2, fill=SOFT, line=None, radius=0.08)
    text(s, x, y + 0.3, 1.95, 0.8, big, size=34, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.12, y + 1.12, 1.71, 1.0, cap, size=18, color=BODY, align=PP_ALIGN.CENTER, spacing=0.95)

# ================================================================== 10 Limitations
s = new_slide(10, "What GymVerse does not solve yet", tag="Evaluation", notes=(
    "Honest limits. Trainer double-booking is checked in the API but not by a database constraint, so two simultaneous requests "
    "can both pass. Check-in blocks suspended members but class enrolment does not. The benchmark uses synthetic data on one PC, "
    "and there has been no multi-user load test. Progress, diet, equipment and feedback tables exist in SQL without screens. "
    "The AI assistant relies on free-tier model quotas."))
limits = [
    "Two admins can double-book a trainer at the same instant",
    "Suspended members can still book classes",
    "Benchmark uses synthetic data on one PC",
    "No multi-user load test yet",
    "Progress, diet and equipment tables have no screens",
    "AI replies depend on free-tier model quotas",
]
for i, item in enumerate(limits):
    x = 0.6 + (i % 3) * 4.15
    y = 1.95 + (i // 3) * 2.35
    box(s, x, y, 3.85, 2.1, fill=WHITE, line=LINE, radius=0.06)
    circle_number(s, x + 0.3, y + 0.3, 0.55, "!", fill=SOFT, color=MUTED, size=20)
    text(s, x + 0.3, y + 0.95, 3.3, 1.1, item, size=20, color=DARK, spacing=0.95)

# ================================================================== 11 Dissemination
s = new_slide(11, "Documented and reproducible in three commands", tag="Dissemination", notes=(
    "Code: the GitHub repository currently holds only the first commit, so the link and QR code are placeholders until the code is pushed. "
    "Documentation: PROJECT_GUIDE.md explains every part of the system; Swagger API docs are served at /api-docs; the SQL files are in database/. "
    "Reproduce: in gymverse-backend run npm install, npm run migrate && npm run seed, npm run dev; then npm install && npm run dev in "
    "gymverse-frontend. node verify_stack.js checks everything on a throwaway database. "
    "Add the licence, any paper or demo link, and talks or publications before presenting."))
rows11 = [
    ("Code", REPO_URL or "[ADD REPO LINK AFTER PUSHING CODE]"),
    ("Documentation", "PROJECT_GUIDE.md · API docs at /api-docs"),
    ("Licence", "[ADD LICENSE]"),
    ("Paper · demo · talks", "[ADD PAPER/DEMO LINK] · [ADD TALKS]"),
]
for i, (lab, val) in enumerate(rows11):
    y = 1.95 + i * 1.18
    text(s, 0.6, y, 6.3, 0.45, lab, size=18, color=MUTED, bold=True)
    placeholder = val.startswith("[")
    text(s, 0.6, y + 0.42, 6.3, 0.7, val, size=20, color=DARK, bold=placeholder)
box(s, 7.3, 1.95, 5.45, 2.95, fill=DARK, line=None, radius=0.05)
text(s, 7.6, 2.12, 4.9, 0.5, "Reproduce (backend)", size=20, color=WHITE, bold=True)
cmds = ["npm install", "npm run migrate && npm run seed", "npm run dev"]
for i, c_ in enumerate(cmds):
    text(s, 7.6, 2.75 + i * 0.66, 0.4, 0.5, str(i + 1), size=18, color=MID, bold=True)
    text(s, 8.0, 2.75 + i * 0.66, 4.6, 0.5, c_, size=18, color=WHITE, font=MONO)
qr_or_placeholder(s, 7.3, 5.1, 1.6, "[ADD QR]")
text(s, 9.15, 5.2, 3.6, 0.5, "Verify on a fresh DB:", size=18, color=MUTED, bold=True)
text(s, 9.15, 5.7, 3.6, 0.5, "node verify_stack.js", size=18, color=DARK, font=MONO)

# ================================================================== 12 Conclusion
s = new_slide(12, "GymVerse keeps gym data correct, private and fast", notes=(
    "Three takeaways: correctness is enforced by the database; privacy is enforced on every record; and measured indexes kept key "
    f"queries fast at scale (up to {best:.0f}× faster). Future work: a database constraint against trainer double-booking, "
    "multi-branch support, and load testing."))
takeaways = [
    "Integrity lives in the database, not the screens",
    "Every record is checked against who is asking",
    f"The right indexes made key queries up to {best:.0f}× faster",
]
for i, t_ in enumerate(takeaways):
    y = 1.95 + i * 1.2
    circle_number(s, 0.6, y, 0.8, str(i + 1), size=24)
    text(s, 1.65, y + 0.1, 10.9, 0.7, t_, size=26, color=DARK, bold=True)
box(s, 0.6, 5.7, 12.1, 0.9, fill=SOFT, line=None, radius=0.1)
text(s, 0.9, 5.83, 11.6, 0.65, "Next: database-level double-booking guard · multi-branch support · load testing",
     size=20, color=BODY, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================== 13 Thank you
s = new_slide(13, "Thank you — questions?", notes="Invite questions. Point to the repository QR code for the code and documentation.")
text(s, 0.6, 1.95, 7.0, 0.5, "Team", size=20, color=MUTED, bold=True)
for i, (name, roll) in enumerate(TEAM):
    text(s, 0.6, 2.5 + i * 0.62, 7.2, 0.55, f"{name}  ·  {roll}", size=20, color=DARK)
text(s, 0.6, 5.2, 7.0, 0.5, "Contact", size=20, color=MUTED, bold=True)
text(s, 0.6, 5.7, 7.2, 0.55, CONTACT_EMAIL or "[ADD EMAIL]", size=22, color=DARK, bold=not CONTACT_EMAIL)
qr_or_placeholder(s, 9.2, 2.0, 3.3, "[ADD QR AFTER PUSHING CODE]")
text(s, 9.2, 5.45, 3.3, 0.5, "Code & documentation", size=18, color=MUTED, align=PP_ALIGN.CENTER)

assert len(prs.slides) == TOTAL_SLIDES, len(prs.slides)
prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides); best speed-up {best:.1f}x; e2e {e2e_passed}/{e2e_total}")
