"""Builds the rubric-only GymVerse deck (presentation_rubrics.pptx) with python-pptx.

Nine slides, one section per evaluation rubric: Novelty, Adaptability, Methodology
Evaluation, and Dissemination & Accessibility. Every number comes from measured results kept
next to this script:
  bench-results.json  index ablation benchmark (synthetic 50k members / 1M check-ins)
  e2e-results.json    live end-to-end run against the real database (116 checks)
Set REPO_URL to replace the repository placeholders with a real link and QR code:
  REPO_URL=https://github.com/... python build_deck_rubrics.py ../presentation_rubrics.pptx
"""
import io
import json
import math
import os
import sys
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
ROOT = HERE.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "presentation_rubrics.pptx"
REPO_URL = os.environ.get("REPO_URL", "").strip()

BENCH = json.loads((HERE / "bench-results.json").read_text())
E2E = json.loads((HERE / "e2e-results.json").read_text())

# ------------------------------------------------------------------ theme
DARK = RGBColor(0x11, 0x18, 0x27)
BODY = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LINE = RGBColor(0xD1, 0xD5, 0xDB)
SOFT = RGBColor(0xF3, 0xF4, 0xF6)
MID = RGBColor(0x9C, 0xA3, 0xAF)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT_TINT = RGBColor(0xDB, 0xEA, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"
SYMBOL = "Segoe UI Symbol"
MONO = "Consolas"
EMOJI = "Segoe UI Emoji"

# One colour per rubric, echoing the evaluation sheet, so each slide maps to its criterion.
RUBRICS = {
    "Novelty": RGBColor(0xDC, 0x26, 0x26),
    "Adaptability": RGBColor(0x2A, 0x7F, 0xA3),
    "Evaluation": RGBColor(0xB0, 0x48, 0x46),
    "Dissemination": RGBColor(0xA1, 0x62, 0x07),
}
TOTAL_SLIDES = 9

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def text(slide, x, y, w, h, lines, size=18, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.0, space_after=0):
    """Text box. `lines` is a string or a list of strings / (string, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    items = [lines] if isinstance(lines, str) else lines
    for i, item in enumerate(items):
        content, opts = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.line_spacing = spacing
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = content
        f = run.font
        f.name = opts.get("font", font)
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=0.08, line_width=1.25, dash=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_width)
        if dash:
            s.line.dash_style = 4
    s.shadow.inherit = False
    return s


def label_box(slide, x, y, w, h, content, fill, line, color, size=18, bold=True, radius=0.1):
    s = box(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = content
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def circle_number(slide, x, y, d, content, fill=ACCENT, color=WHITE, size=22, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(2)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = content
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
    return s


def new_slide(number, title, rubric=None, notes=""):
    """Content slide. The title bar and the tag take the rubric's colour."""
    slide = prs.slides.add_slide(BLANK)
    colour = RUBRICS.get(rubric, ACCENT)
    if title:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.38), Inches(0.55), Inches(0.08), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = colour
        bar.line.fill.background()
        bar.shadow.inherit = False
        text(slide, 0.6, 0.42, 10.3, 1.25, title, size=34, color=DARK, bold=True, spacing=0.95)
    if rubric:
        label_box(slide, 11.0, 0.5, 1.8, 0.46, rubric, fill=WHITE, line=colour, color=colour, size=14, radius=0.5)
    if number:
        text(slide, 12.2, 6.95, 0.7, 0.4, str(number), size=14, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def qr_or_placeholder(slide, x, y, size, placeholder):
    if REPO_URL:
        code = qrcode.QRCode(border=1, box_size=10)
        code.add_data(REPO_URL)
        code.make(fit=True)
        buf = io.BytesIO()
        code.make_image(fill_color="#111827", back_color="white").save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(size), Inches(size))
        return
    s = box(slide, x, y, size, size, fill=SOFT, line=MID, dash=True, radius=0.04)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = placeholder
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(16), True, DARK


def style_chart(chart):
    chart.font.name = FONT
    chart.font.size = Pt(18)
    chart.font.color.rgb = BODY
    # No automatic title (a single-series chart otherwise shows its series name), and the
    # category axis sits at the value axis minimum so bars below 1 on a log scale grow upwards.
    chart.has_title = False
    chart.value_axis.crosses = XL_AXIS_CROSSES.MINIMUM


def table(slide, x, y, w, col_widths, row_heights):
    gf = slide.shapes.add_table(len(row_heights), len(col_widths), Inches(x), Inches(y), Inches(w),
                                Inches(sum(row_heights)))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for c, cw in enumerate(col_widths):
        tbl.columns[c].width = Inches(cw)
    for r, rh in enumerate(row_heights):
        tbl.rows[r].height = Inches(rh)
    return tbl


def cell(tbl, r, c, content, size=18, bold=False, color=BODY, fill=WHITE, align=PP_ALIGN.CENTER, font=FONT):
    ce = tbl.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    ce.margin_left = ce.margin_right = Inches(0.1)
    tf = ce.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = content
    run.font.name, run.font.size, run.font.bold, run.font.color.rgb = font, Pt(size), bold, color


# ------------------------------------------------------------------ data
bench_q = {q["key"]: q for q in BENCH["queries"]}
bench_order = [("payments_member", "Payment history"), ("subscriptions_list", "Subscriptions list"),
               ("attendance_log", "Attendance log"), ("member_search", "Member search")]
median = lambda key, setup: bench_q[key]["results"][setup]["median_ms"]
speedups = {k: median(k, "before16") / median(k, "after16") for k, _ in bench_order}
best = max(speedups.values())
slowest_optimised = max(median(k, "after16") for k, _ in bench_order)
vol = BENCH["volumes"]

groups = [
    ("Accounts & security", ["Login", "Register member", "Trainer approval", "Account security", "Access control"]),
    ("Training & classes", ["Exercises", "Workouts", "Classes"]),
    ("Dashboards, health & UI", ["Dashboards & reports", "Health", "Frontend"]),
    ("Members & plans", ["Members", "Plans"]),
    ("Payments & subscriptions", ["Subscriptions & payments"]),
    ("Attendance & rosters", ["Attendance", "Trainer rosters"]),
]
area_pass = {}
for r in E2E["results"]:
    area_pass.setdefault(r["area"], [0, 0])
    area_pass[r["area"]][0 if r["ok"] else 1] += 1
group_counts = [(name, sum(area_pass[a][0] for a in areas), sum(sum(area_pass[a]) for a in areas))
                for name, areas in groups]
e2e_passed = sum(1 for r in E2E["results"] if r["ok"])
e2e_total = len(E2E["results"])
assert sum(t for _, _, t in group_counts) == e2e_total, "every end-to-end check must be in a group"

TEAM = [("NYASHADAM SRIKAR", "2510030006"), ("MALLALA SRI CHARAN", "2510030246"),
        ("YOGENDRA UNGARALA", "2510030147"), ("GORANTAL PRANAV", "2510030008")]

# ================================================================== 1 Title + rubric map
s = new_slide(None, None, notes=(
    "GymVerse: a PostgreSQL-backed gym management system with a web app and an AI assistant. "
    "This deck is organised by the four evaluation rubrics; the coloured tag on each slide names its rubric."))
text(s, 0.8, 0.8, 7.4, 0.5, "DATABASE MANAGEMENT SYSTEMS · MINI PROJECT", size=18, color=ACCENT, bold=True)
text(s, 0.8, 1.3, 7.4, 1.3, "GymVerse", size=66, color=DARK, bold=True)
text(s, 0.8, 2.65, 7.4, 1.1, "The Digital Transformation of Fitness Management", size=28, color=DARK)
agenda = [("Novelty", "Novelty of developed methodology", "2–3"),
          ("Adaptability", "Adaptability of methodology", "4"),
          ("Evaluation", "Methodology evaluation", "5–7"),
          ("Dissemination", "Dissemination & accessibility", "8–9")]
for i, (key, label, slides) in enumerate(agenda):
    y = 3.85 + i * 0.58
    box(s, 0.8, y + 0.1, 0.22, 0.22, fill=RUBRICS[key], line=None, radius=0.25)
    text(s, 1.2, y, 5.2, 0.45, label, size=18, color=DARK)
    text(s, 6.4, y, 1.4, 0.45, f"slide {slides}", size=18, color=MUTED, align=PP_ALIGN.RIGHT)
box(s, 8.35, 0.8, 4.4, 5.1, fill=SOFT, line=None, radius=0.05)
text(s, 8.7, 1.05, 3.8, 0.5, "Team", size=20, color=DARK, bold=True)
for i, (name, roll) in enumerate(TEAM):
    y = 1.7 + i * 1.02
    text(s, 8.7, y, 3.9, 0.45, name, size=18, color=DARK, bold=True)
    text(s, 8.7, y + 0.4, 3.9, 0.45, roll, size=18, color=MUTED)
rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.25), Inches(11.95), Inches(0.03))
rule.fill.solid(); rule.fill.fore_color.rgb = LINE; rule.line.fill.background(); rule.shadow.inherit = False
text(s, 0.8, 6.4, 11.9, 0.5, "Department of CSE  ·  KLH  ·  Faculty Guide: P LALITHA", size=18, color=MUTED)

# ================================================================== 2 Novelty: comparison
s = new_slide(2, "Only GymVerse enforces gym rules inside the database", rubric="Novelty", notes=(
    "Rows are approaches; columns are capabilities. 'Partial' for spreadsheets: data validation, sheet-level sharing "
    "and version history exist but do not enforce gym rules or per-record access. "
    "[ADD APPROACH]: add another existing system you compared against, and fill in its row."))
cols = ["Approach", "One source of truth", "Rules enforced in database", "Safe with concurrent users",
        "Per-record access", "Audit trail", "AI on live data"]
rows = [
    ("Paper registers", ["✕", "✕", "✕", "✕", "✕", "✕"]),
    ("Spreadsheets", ["✕", "Partial", "✕", "Partial", "Partial", "✕"]),
    ("[ADD APPROACH]", ["[ADD]"] * 6),
    ("GymVerse (ours)", ["✓"] * 6),
]
tbl = table(s, 0.6, 1.9, 12.1, [2.95] + [(12.1 - 2.95) / 6] * 6, [1.4] + [0.8] * len(rows))
for c, h in enumerate(cols):
    cell(tbl, 0, c, h, bold=True, color=DARK, fill=SOFT, align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
for r, (name, marks) in enumerate(rows, start=1):
    ours = name.startswith("GymVerse")
    fill = ACCENT_TINT if ours else WHITE
    cell(tbl, r, 0, name, size=20, bold=True, color=ACCENT if ours else DARK, fill=fill, align=PP_ALIGN.LEFT)
    for c, m in enumerate(marks, start=1):
        symbol = m in ("✓", "✕")
        cell(tbl, r, c, m, size=26 if symbol else 18, bold=ours or symbol,
             color=ACCENT if ours else (MID if m == "✕" else BODY), fill=fill,
             font=SYMBOL if symbol else FONT)

# ================================================================== 3 Novelty: contributions in the architecture
s = new_slide(3, "Three contributions, built into every request", rubric="Novelty", notes=(
    "Every request passes rate limiting, a session check (JWT plus server-side revocation), role and per-record access, "
    "and input validation; controllers run multi-row writes in transactions. The numbered badges mark our contributions. "
    "1) Rules live in the database: partial unique indexes (one active subscription, one open check-in), row locks for "
    "bookings and subscriptions, triggers for capacity and auditing. "
    "2) Four roles checked per record: a member reaches only their own data, a trainer only their roster. "
    "3) The AI assistant reads the live plan catalogue and falls back across Gemini models, then Groq. "
    "Chat history in MongoDB is optional."))
label_box(s, 0.5, 2.05, 1.8, 1.0, "React app", fill=SOFT, line=LINE, color=BODY)
arrow(s, 2.32, 2.55, 2.68, 2.55)
guards = [("Rate limit", False), ("Session check", True), ("Role + record access", True), ("Input validation", False)]
gx = [2.7, 4.55, 6.4, 8.25]
for i, ((name, novel), x) in enumerate(zip(guards, gx)):
    label_box(s, x, 2.05, 1.6, 1.0, name,
              fill=ACCENT if novel else SOFT, line=ACCENT if novel else LINE,
              color=WHITE if novel else BODY)
    if i < 3:
        arrow(s, x + 1.62, 2.55, gx[i + 1] - 0.02, 2.55)
text(s, 2.7, 3.12, 3.6, 0.45, "Express API: four guards", size=18, color=MUTED)
arrow(s, 9.05, 3.07, 9.05, 3.73)
label_box(s, 6.4, 3.75, 3.45, 0.85, "Controllers + transactions", fill=SOFT, line=LINE, color=BODY)
box(s, 10.25, 1.75, 2.6, 4.65, fill=WHITE, line=DARK, radius=0.05, line_width=1.5)
text(s, 10.25, 1.82, 2.6, 0.5, "PostgreSQL", size=20, color=DARK, bold=True, align=PP_ALIGN.CENTER)
chips = [("Unique & partial indexes", True), ("Row locks", True), ("Audit + capacity triggers", True),
         ("Views & functions", False)]
for i, (name, novel) in enumerate(chips):
    label_box(s, 10.4, 2.4 + i * 0.97, 2.3, 0.8, name,
              fill=ACCENT if novel else SOFT, line=ACCENT if novel else LINE,
              color=WHITE if novel else BODY)
arrow(s, 9.87, 4.17, 10.23, 4.17)
label_box(s, 6.4, 5.25, 1.65, 0.9, "AI model fallback", fill=ACCENT, line=ACCENT, color=WHITE)
label_box(s, 8.2, 5.25, 1.65, 0.9, "Chat history (MongoDB)", fill=SOFT, line=LINE, color=BODY)
arrow(s, 7.22, 4.62, 7.22, 5.23)
arrow(s, 9.02, 4.62, 9.02, 5.23)
label_box(s, 4.2, 5.25, 1.95, 0.9, "Gemini → Groq", fill=SOFT, line=LINE, color=BODY)
arrow(s, 6.38, 5.7, 6.17, 5.7)
for num, bx, by in [("1", 12.45, 2.2), ("2", 7.78, 1.84), ("3", 7.83, 5.04)]:
    circle_number(s, bx, by, 0.42, num, fill=WHITE, color=ACCENT, size=16, line=ACCENT)
box(s, 0.5, 3.65, 3.5, 2.1, fill=ACCENT_TINT, line=None, radius=0.06)
text(s, 0.75, 3.78, 3.1, 0.5, "Key contributions", size=20, color=ACCENT, bold=True)
for i, item in enumerate(["Rules in the database", "Per-record access", "AI with model fallback"]):
    y = 4.35 + i * 0.45
    text(s, 0.75, y, 0.35, 0.45, str(i + 1), size=18, color=ACCENT, bold=True)
    text(s, 1.1, y, 2.8, 0.45, item, size=18, color=DARK)
box(s, 0.6, 6.53, 0.3, 0.3, fill=ACCENT, line=None, radius=0.2)
text(s, 1.0, 6.45, 2.8, 0.45, "Novel in GymVerse", size=18, color=BODY)
box(s, 3.8, 6.53, 0.3, 0.3, fill=SOFT, line=LINE, radius=0.2)
text(s, 4.2, 6.45, 2.8, 0.45, "Standard component", size=18, color=BODY)

# ================================================================== 4 Adaptability
s = new_slide(4, "One schema pattern fits any booking business", rubric="Adaptability", notes=(
    "The method models generic concepts rather than gym-only ones: people, plans and subscriptions, sessions with "
    "capacity and bookings, a check-in log, and staff who see only their own clients. The table shows how each maps "
    "onto other businesses by renaming. Deployment scales from one PC (AI assistant, email and chat history can all be "
    "switched off) to a hosted server, where one Node process can serve both the API and the built web app. "
    f"The benchmark shows every optimised key query stays under {math.ceil(slowest_optimised)} ms at "
    f"{vol['attendance']:,} check-ins. Plans, AI models, email, chat history and production mode are configuration; "
    "integrity rules, record access and the audit trail stay the same everywhere."))
map_cols = ["Reusable pattern", "GymVerse", "Yoga studio", "Physio clinic", "Co-working"]
map_rows = [
    ("Client", ["Member", "Student", "Patient", "Member"]),
    ("Plan → subscription", ["Membership plan", "Class pass", "Care package", "Desk plan"]),
    ("Session with capacity", ["Class schedule", "Yoga class", "Appointment", "Room slot"]),
    ("Staff see own clients", ["Trainer roster", "Own classes", "Own patients", "Own bookings"]),
]
tbl = table(s, 0.6, 1.85, 12.1, [3.1] + [(12.1 - 3.1) / 4] * 4, [0.6] * (len(map_rows) + 1))
for c, h in enumerate(map_cols):
    cell(tbl, 0, c, h, size=17, bold=True, color=ACCENT if c == 1 else DARK, fill=ACCENT_TINT if c == 1 else SOFT,
         align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
for r, (pattern, values) in enumerate(map_rows, start=1):
    cell(tbl, r, 0, pattern, size=17, bold=True, color=DARK, align=PP_ALIGN.LEFT)
    for c, v in enumerate(values, start=1):
        cell(tbl, r, c, v, size=17, color=ACCENT if c == 1 else BODY, bold=c == 1,
             fill=ACCENT_TINT if c == 1 else WHITE)
tiles4 = [
    ("💻 → ☁", "One PC to hosted server", "AI, email, chat history optional"),
    # Round up: "< 8 ms" would be false for an 8.4 ms median.
    (f"< {math.ceil(slowest_optimised)} ms", f"at {vol['attendance'] // 1_000_000}M check-ins", "every optimised key query"),
    ("⚙", "Configure, don't recode", "Plans · AI models · email"),
]
for i, (big, head, cap) in enumerate(tiles4):
    x = 0.6 + i * 4.1
    box(s, x, 5.15, 3.9, 1.6, fill=SOFT, line=None, radius=0.08)
    emoji = not big[0].isdigit() and not big.startswith("<")
    text(s, x + 0.25, 5.28, 1.6, 0.7, big, size=26 if emoji else 24, color=ACCENT, bold=True,
         font=EMOJI if emoji else FONT)
    text(s, x + 0.25, 5.9, 3.5, 0.45, head, size=18, color=DARK, bold=True)
    text(s, x + 0.25, 6.3, 3.5, 0.45, cap, size=16, color=MUTED)

# ================================================================== 5 Evaluation: how we evaluated
s = new_slide(5, "Three independent methods test speed and correctness", rubric="Evaluation", notes=(
    "1) Performance: the app's own list and search queries, run against three index setups (no indexes, the original "
    f"indexes, and our optimised indexes) on a synthetic database of {vol['members']:,} members and {vol['attendance']:,} "
    f"check-ins; each ran {BENCH['runs']} times under EXPLAIN ANALYZE and we report the median. "
    f"2) Live end-to-end: {e2e_total} checks drove every feature through the real API as each role, and every write was "
    "confirmed with a direct database query. "
    "3) Automated tests: Jest unit and route tests, Vitest frontend tests, and a stack check that builds a fresh database "
    "from the SQL scripts. Raw results are kept in bench-results.json and e2e-results.json."))
methods = [
    ("Performance ablation", "Question", "Do our indexes make real queries faster?",
     "Method", f"App queries × 3 index setups · {vol['members'] // 1000}k members, "
               f"{vol['attendance'] // 1_000_000}M check-ins · {BENCH['runs']} runs each",
     "Measure", "Median time (EXPLAIN ANALYZE)"),
    ("Live end-to-end run", "Question", "Does every feature work for every role?",
     "Method", "Drive the real API as admin, receptionist, trainer and member",
     "Measure", f"{e2e_total} checks; each write confirmed in the database"),
    ("Automated test suites", "Question", "Do the parts stay correct as code changes?",
     "Method", "Unit & route tests · frontend tests · fresh-database stack check",
     "Measure", "Tests passed; server errors"),
]
for i, (head, *pairs) in enumerate(methods):
    x = 0.6 + i * 4.1
    box(s, x, 1.9, 3.9, 4.85, fill=WHITE, line=LINE, radius=0.05)
    circle_number(s, x + 0.25, 2.1, 0.6, str(i + 1), fill=RUBRICS["Evaluation"], size=20)
    text(s, x + 1.0, 2.1, 2.8, 0.6, head, size=20, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    for j, y in enumerate((2.95, 4.05, 5.5)):   # the method text runs to three lines
        label, value = pairs[2 * j], pairs[2 * j + 1]
        text(s, x + 0.25, y, 3.4, 0.4, label.upper(), size=14, color=MUTED, bold=True)
        text(s, x + 0.25, y + 0.35, 3.45, 0.9, value, size=17, color=BODY, spacing=0.95)

# ================================================================== 6 Evaluation: benchmark (ablation)
s = new_slide(6, f"Optimised indexes make key queries up to {best:.0f}× faster", rubric="Evaluation", notes=(
    f"Setup: a throwaway PostgreSQL {BENCH.get('postgres', '')} database with synthetic volumes "
    f"({vol['members']:,} members, {vol['subscriptions']:,} subscriptions, {vol['payments']:,} payments, "
    f"{vol['attendance']:,} check-ins) on one Windows PC. Each query ran {BENCH['runs']} times per index setup; "
    "medians on a log scale. No indexes / original indexes / optimised indexes: "
    + " ".join(f"{label}: {median(k, 'none'):.1f} / {median(k, 'before16'):.1f} / {median(k, 'after16'):.2f} ms."
               for k, label in bench_order)
    + " The original indexes and no indexes both used a full scan for the attendance log, so that gap is timing noise."))
cd = CategoryChartData()
cd.categories = [label for _, label in bench_order]
cd.add_series("No indexes", [median(k, "none") for k, _ in bench_order])
cd.add_series("Original indexes", [median(k, "before16") for k, _ in bench_order])
cd.add_series("Optimised (GymVerse)", [median(k, "after16") for k, _ in bench_order])
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.8), Inches(8.9), Inches(5.05), cd).chart
style_chart(chart)
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.include_in_layout = False
chart.legend.font.size = Pt(18)
plot = chart.plots[0]
plot.gap_width = 60
plot.overlap = -10
for ser, color in zip(plot.series, [LINE, MID, ACCENT]):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = color
    ser.format.line.fill.background()
ours = plot.series[2]
ours.data_labels.show_value = True
ours.data_labels.number_format = '0.0" ms"'
ours.data_labels.number_format_is_linked = False
ours.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ours.data_labels.font.size = Pt(16)
ours.data_labels.font.bold = True
ours.data_labels.font.color.rgb = ACCENT
va = chart.value_axis
scaling = va._element.find(qn("c:scaling"))
scaling.insert(0, scaling.makeelement(qn("c:logBase"), {"val": "10"}))
va.minimum_scale = 0.1
va.maximum_scale = 1000
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = SOFT
va.format.line.fill.background()
va.tick_labels.font.size = Pt(16)
va.tick_labels.font.color.rgb = MUTED
va.has_title = True
va.axis_title.text_frame.text = "Median time (ms, log scale)"
va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
va.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
chart.category_axis.tick_labels.font.size = Pt(18)
chart.category_axis.format.line.color.rgb = LINE
text(s, 9.55, 1.95, 3.5, 0.5, "Speed-up vs original", size=18, color=MUTED, bold=True)
for i, (k, label) in enumerate(sorted(bench_order, key=lambda kv: -speedups[kv[0]])):
    y = 2.5 + i * 0.9
    text(s, 9.55, y, 1.45, 0.7, f"{speedups[k]:.0f}×" if speedups[k] >= 10 else f"{speedups[k]:.1f}×",
         size=32, color=ACCENT, bold=True)
    text(s, 10.95, y + 0.12, 2.2, 0.6, label, size=18, color=BODY)
text(s, 9.55, 6.0, 3.6, 0.85,
     [f"{vol['members'] // 1000}k members, {vol['attendance'] // 1_000_000}M check-ins",
      f"median of {BENCH['runs']} runs"], size=18, color=MUTED)

# ================================================================== 7 Evaluation: functional results
s = new_slide(7, f"All {e2e_passed} end-to-end checks pass across every role", rubric="Evaluation", notes=(
    f"Live run against the real database: {e2e_passed} of {e2e_total} checks, each write confirmed with a direct "
    "database query. Unit and route tests: 235 passed, 1 skipped. The stack check builds a fresh database from the "
    "SQL scripts and passes 71 of 71. Frontend tests: 7 of 7. No 5xx responses during the run."))
cd = CategoryChartData()
cd.categories = [g for g, _, _ in group_counts]
cd.add_series("Checks passed", [p for _, p, _ in group_counts])
chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.85), Inches(7.9), Inches(4.95), cd).chart
style_chart(chart)
chart.has_legend = False
plot = chart.plots[0]
plot.gap_width = 55
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = ACCENT
ser.format.line.fill.background()
ser.data_labels.show_value = True
ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser.data_labels.font.size = Pt(18)
ser.data_labels.font.bold = True
ser.data_labels.font.color.rgb = DARK
ca = chart.category_axis
ca.reverse_order = True
ca.tick_labels.font.size = Pt(18)
ca.format.line.color.rgb = LINE
va = chart.value_axis
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = max(t for _, _, t in group_counts) * 1.2
va.minimum_scale = 0
tiles = [("235", "unit & route tests pass"), ("71/71", "stack checks on a fresh database"),
         ("7/7", "frontend tests"), ("0", "server errors (5xx)")]
for i, (big, cap) in enumerate(tiles):
    x = 8.75 + (i % 2) * 2.1
    y = 1.95 + (i // 2) * 2.45
    box(s, x, y, 1.95, 2.2, fill=SOFT, line=None, radius=0.08)
    text(s, x, y + 0.3, 1.95, 0.8, big, size=34, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.12, y + 1.12, 1.71, 1.0, cap, size=18, color=BODY, align=PP_ALIGN.CENTER, spacing=0.95)

# ================================================================== 8 Dissemination
s = new_slide(8, "Everything needed to study or rebuild GymVerse is shared", rubric="Dissemination", notes=(
    "Code: the repository link and QR code are placeholders until the code is pushed. "
    "Documentation: PROJECT_GUIDE.md explains every part of the system; the backend README and CHAT_HISTORY.md cover "
    "setup and the optional MongoDB history; Swagger API docs are served at /api-docs. "
    "Database design: the ER model opens in TerraER (gymverse_er_model.xml) and is exported as a PNG; "
    "build_database.ps1 / build_database.sh create the whole database with sample data in one command. "
    "Reproduce: build the database, copy .env.example to .env in gymverse-backend, then npm install and npm run dev in "
    "the backend and frontend. Add the licence and any paper, demo or talk links before presenting."))
artefacts = [
    ("Code", REPO_URL or "[ADD REPO LINK]"),
    ("Documentation", "PROJECT_GUIDE.md · README · /api-docs"),
    ("Database design", "ER model (29 entities, opens in TerraER)"),
    ("Build", "One-command database script · sample data"),
    ("Licence · paper · demo", "[ADD LICENSE] · [ADD PAPER/DEMO LINK]"),
]
for i, (lab, val) in enumerate(artefacts):
    y = 1.9 + i * 0.97
    text(s, 0.6, y, 6.3, 0.4, lab, size=16, color=MUTED, bold=True)
    text(s, 0.6, y + 0.36, 6.3, 0.55, val, size=19, color=DARK, bold=val.startswith("["))
box(s, 7.3, 1.9, 5.45, 3.05, fill=DARK, line=None, radius=0.05)
text(s, 7.6, 2.07, 4.9, 0.5, "Rebuild in three steps", size=20, color=WHITE, bold=True)
steps = [("Database + sample data", ".\\build_database.ps1"),
         ("API  (gymverse-backend)", "npm install; npm run dev"),
         ("Web app  (gymverse-frontend)", "npm install; npm run dev")]
for i, (what, cmd) in enumerate(steps):
    y = 2.65 + i * 0.75
    text(s, 7.6, y, 0.4, 0.5, str(i + 1), size=18, color=MID, bold=True)
    text(s, 8.0, y - 0.04, 4.6, 0.4, what, size=14, color=MID)
    text(s, 8.0, y + 0.26, 4.6, 0.45, cmd, size=17, color=WHITE, font=MONO)
qr_or_placeholder(s, 7.3, 5.2, 1.6, "[ADD QR]")
text(s, 9.15, 5.3, 3.6, 0.5, "Check a fresh build:", size=18, color=MUTED, bold=True)
text(s, 9.15, 5.8, 3.6, 0.5, "node verify_stack.js", size=18, color=DARK, font=MONO)

# ================================================================== 9 Accessibility
s = new_slide(9, "Usable from any browser, by every role", rubric="Dissemination", notes=(
    "For people at the gym: nothing to install beyond a web browser; four role-specific views (admin, receptionist, "
    "trainer, member) show each person only what they need; accounts are self-service with email verification and "
    "forgotten-password reset; signed-in users can ask the AI assistant in plain language. Interactive controls carry "
    "ARIA labels and roles, and Escape closes dialogs. "
    "For developers: interactive API docs at /api-docs; one Node server can host both the API and the built web app; "
    "the optional services (AI, email, MongoDB chat history) can be left out and the core still runs."))
columns = [
    ("For gym staff & members", [
        ("🌐", "Any web browser", "nothing to install"),
        ("👥", "Four role-specific views", "admin · receptionist · trainer · member"),
        ("🔑", "Self-service accounts", "email verification · password reset"),
        ("💬", "AI assistant", "ask questions in plain language"),
    ]),
    ("For developers & evaluators", [
        ("📘", "Interactive API docs", "Swagger UI at /api-docs"),
        ("⌨", "Keyboard & screen readers", "ARIA labels on controls · Escape closes dialogs"),
        ("🖥", "One server", "hosts the API and the built web app"),
        ("🧩", "Optional services", "runs without AI, email or MongoDB"),
    ]),
]
for ci, (head, items) in enumerate(columns):
    x = 0.6 + ci * 6.15
    text(s, x, 1.85, 5.95, 0.5, head, size=20, color=RUBRICS["Dissemination"], bold=True)
    for i, (icon, title_, cap) in enumerate(items):
        y = 2.45 + i * 1.08
        box(s, x, y, 5.95, 0.95, fill=SOFT if ci else WHITE, line=None if ci else LINE, radius=0.12)
        text(s, x + 0.2, y + 0.14, 0.7, 0.7, icon, size=24, font=EMOJI, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.95, y + 0.08, 4.9, 0.45, title_, size=18, color=DARK, bold=True)
        text(s, x + 0.95, y + 0.48, 4.9, 0.42, cap, size=15, color=MUTED)
text(s, 0.6, 6.85, 11.0, 0.45, "Thank you  ·  Questions?", size=18, color=ACCENT, bold=True)

assert len(prs.slides) == TOTAL_SLIDES, len(prs.slides)
prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides); best speed-up {best:.1f}x; e2e {e2e_passed}/{e2e_total}")
