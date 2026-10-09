from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
SOURCE = Path('C:/Users/John/Downloads/Telegram Desktop/PBL Final Documentation templet (2).docx')
OUT = ROOT / 'output/GymVerse PBL Final Report.docx'
doc = Document(SOURCE)
original = list(doc.paragraphs)
section_xml = [deepcopy(s._sectPr) for s in doc.sections]
source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
with ZipFile(SOURCE) as z:
    inventory = {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
(QA/'source_inventory.json').write_text(json.dumps(inventory, indent=2))
(QA/'artifact.md').write_text(f'''# GymVerse report template contract
Reference: {SOURCE}
SHA256: {source_hash}
Reference rendering: reference.pdf and reference/page-1.png through page-8.png; reviewed all pages.
Seven A4 portrait sections. Cover margins 0.62in horizontal, 0.31in top, 1in bottom. Body 0.63in horizontal, 0.75in top, 1in bottom.
Cover uses centered Times New Roman, 20pt title, 13pt submission text, 14pt degree and student names, 11.5pt guide and institution. Preserve KLH logo in paragraph 18 and original institution wording.
Body uses source Body Text, Heading 1 and Heading 2 roles, Times New Roman theme and 1.5 line spacing. Preserve named styles and heading numbering. Source title role 24pt.
Slots: document.xml paragraphs 1,2,7,8,11-13,16,20,23 fill with GymVerse and supplied details. Clone student row twice. Paragraphs 24-28 retain front matter headings and receive project content. Paragraphs 29 onward replace all sample medical prose and illustrations with project report. Preserve original section geometries and seven section definitions; move section endings after their new content. New chapter page breaks permit a longer report.
Preserve-only: all original package parts other than document.xml, document relationships, and content types. New media may be added. Source styles, numbering, theme, institution media, headers, footers and settings are byte-preserved by final package merge.
New content: two project diagrams and four tables replace sample figures. Static contents and lists are verified against final Word PDF page locations. No fabricated signatures, course code, live deployment results, or performance measurements.
Academic year assumption: 2026-2027; year and semester: II Year II Semester. Course code omitted pending user information.
Rendering deviation: packaged LibreOffice unavailable; installed Word exports PDF for Poppler PNG review.
''', encoding='utf-8')

def replace(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]: r.text = ''
    else: p.add_run(text)

replace(original[1], 'GymVerse Gym Management System')
replace(original[2], 'A Project Based Learning Report Submitted in partial fulfilment of the requirements for the award of the degree')
replace(original[7], 'in the Department of Computer Science and Engineering')
replace(original[8], 'II Year, II Semester')
for r in original[8].runs: r.font.color.rgb = RGBColor(0,0,0)
students = [('2510030006','N Srikar'),('2510030008','G Pranav'),('2510030147','Yogendra'),('2510030246','Sri Charan'),('2510030306','K Abhineshwar')]
for p, (roll,name) in zip(original[11:14], students): replace(p, f'{roll}: {name}')
anchor = original[13]._p
from docx.text.paragraph import Paragraph
for roll,name in students[3:]:
    el = deepcopy(original[13]._p); anchor.addnext(el); anchor=el
    replace(Paragraph(el, doc._body), f'{roll}: {name}')
replace(original[16], 'P Lalitha')
replace(original[20], 'Department of Computer Science and Engineering')
replace(original[23], 'Academic Year 2026-2027')
for i in [9,14,17,19]:
    original[i].paragraph_format.space_before=Pt(0)
    original[i].paragraph_format.space_after=Pt(0)
    original[i].paragraph_format.line_spacing=1

for p in original[24:]: p._p.getparent().remove(p._p)

def body(text):
    listing = re.match(r'^(.*?)\s*\.{3,}\s*(\d+)$', text)
    if listing:
        label = listing.group(1).rstrip()
        for arabic, roman in [('1  ', 'I  '), ('2  ', 'II  '), ('3  ', 'III  '), ('4  ', 'IV  '), ('5  ', 'V  ')]:
            if label.startswith(arabic): label = roman + label[len(arabic):]
        text = label + '\t' + listing.group(2)
    p=doc.add_paragraph(style='Body Text')
    if listing:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.first_line_indent=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(11)
    return p

def heading(text,level=1,newpage=False):
    p=doc.add_paragraph(text, style=f'Heading {level}')
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.page_break_before=newpage
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.space_before=Pt(10)
    p.paragraph_format.space_after=Pt(6)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(12 if level==1 else 11); r.font.color.rgb=RGBColor(0,0,0); r.bold=level==1
    return p

def front(text):
    p=doc.add_paragraph(text)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(20)
    r=p.runs[0]; r.font.name='Times New Roman'; r.font.size=Pt(15); r.bold=True

def section_end(index):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1
    p._p.get_or_add_pPr().append(deepcopy(section_xml[index]))

def caption(text):
    p=body(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1
    for r in p.runs: r.italic=True; r.font.size=Pt(10)

def table(number,title,headers,rows,widths):
    p=body(f'Table {number} {title}'); p.paragraph_format.keep_with_next=True
    for r in p.runs: r.bold=True
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers): c.text=h
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=v
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci])
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6'); pr.append(sh)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
                for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(10); r.bold=ri==0
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def diagram(filename, nodes, edges, size):
    im=Image.new('RGB',size,'white'); draw=ImageDraw.Draw(im)
    f=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',24)
    for a,b,label in edges:
        draw.line([a,b],fill='#555555',width=3)
        x,y=b; draw.polygon([(x,y),(x-7,y-12),(x+7,y-12)],fill='#555555') if a[0]==b[0] else None
        if label:
            label = label.replace('1 : many','1 : N').replace('many : 1','N : 1')
            x=(a[0]+b[0])/2; y=(a[1]+b[1])/2
            if a[1]==b[1]: x-=draw.textlength(label,font=f)/2; y-=34
            else: x+=10; y-=18
            draw.text((x,y),label,font=f,fill='black')
    for box,text in nodes:
        draw.rounded_rectangle(box,radius=10,fill='#f0f2f4',outline='#333333',width=3)
        bb=draw.multiline_textbbox((0,0),text,font=f,spacing=8,align='center')
        x=(box[0]+box[2]-bb[2])/2; y=(box[1]+box[3]-bb[3])/2
        draw.multiline_text((x,y),text,font=f,fill='black',spacing=8,align='center')
    im.save(QA/filename)

diagram('architecture.png',[
 ((330,15,870,105),'Browser\nReact interface'),
 ((270,175,930,280),'Render backend\nExpress API and authorization'),
 ((20,390,390,495),'Neon PostgreSQL\nGym records and accounts'),
 ((440,390,800,495),'Optional MongoDB\nChat conversations'),
 ((850,390,1190,495),'AI providers and SMTP\nReplies and email')],
 [((600,105),(600,175),'HTTPS JSON'),((600,280),(200,390),''),((600,280),(620,390),''),((600,280),(1020,390),'')],(1220,520))
diagram('relationships.png',[
 ((20,20,310,100),'Members'),((460,20,750,100),'Subscriptions'),((900,20,1190,100),'Membership plans'),
 ((460,210,750,290),'Payments'),((20,400,310,480),'Class bookings'),((460,400,750,480),'Class schedules'),((900,400,1190,480),'Fitness classes'),((900,210,1190,290),'Trainers')],
 [((310,60),(460,60),'1 : many'),((750,60),(900,60),'many : 1'),((605,100),(605,210),'1 : many'),((165,100),(165,400),'1 : many'),((310,440),(460,440),'many : 1'),((750,440),(900,440),'many : 1'),((1045,290),(750,400),'1 : many')],(1220,520))

front('Abstract')
body('GymVerse is a web based gym management system that brings membership administration, attendance, trainer assignments, workouts and class bookings into a common application. We developed the project to connect routine gym operations to a structured database, so that staff and members can work with consistent records rather than separate registers or spreadsheets.')
body('The application uses a React frontend and an Express backend. PostgreSQL stores accounts and operational records, with Neon serving as the hosted database and Render as the backend hosting platform. The design separates membership plans from individual subscriptions, scheduled classes from bookings, and exercise definitions from assigned workout plans. Foreign keys, unique constraints and transactions protect relationships between these records. Role checks and record ownership checks restrict access for administrators, receptionists, trainers and members.')
body('The public website presents training information, membership plans and upcoming classes. Authenticated pages support the operational workflows. An optional assistant answers questions using the plan catalogue; optional MongoDB storage retains chat conversations without becoming the main gym database. Password recovery and email confirmation depend on an email service configuration.')
body('Recorded project verification completed 71 full-stack checks, 16 frontend tests and three development-origin tests, and produced a successful frontend build. These results demonstrate the tested local flows; they do not establish production capacity or availability. The project illustrates how relational modelling, API validation and a role-aware interface can support a working gym administration application. Further work includes production monitoring, broader concurrency testing and payment gateway integration.')
body('Keywords: gym management, PostgreSQL, React, Express, role based access, database transactions.')
section_end(1)
front('List of Figures')
body('Figure 1  GymVerse application architecture ........................................ 7')
body('Figure 2  Core membership and class relationships ............................... 8')
section_end(2)
section_end(3)
front('List of Tables')
body('Table 1  User roles and responsibilities ................................................ 6')
body('Table 2  Core database entities ............................................................ 8')
body('Table 3  Functional verification scenarios ........................................... 11')
body('Table 4  Recorded verification results ................................................ 12')
section_end(4)
front('Table of Contents')
for text in ['Abstract ................................................................................................ 2','List of Figures ...................................................................................... 3','List of Tables ....................................................................................... 4','1  Introduction ..................................................................................... 6','2  Methodology .................................................................................... 7','    Architecture and technology choices ............................................... 7','    Database design ............................................................................... 8','    Integrity and security ....................................................................... 9','    Implementation and deployment ................................................... 10','3  Experiments ................................................................................... 11','4  Results ........................................................................................... 12','5  Conclusion and Future Work ......................................................... 13','References ......................................................................................... 14']:
    p=body(text); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
section_end(5)

p=doc.add_paragraph('GymVerse Gym Management System',style='paper title')
p.paragraph_format.space_after=Pt(10)
heading('Introduction')
body('A gym manages several related activities each day: registering members, assigning subscriptions, recording payments, tracking attendance and coordinating training. When these activities are maintained in separate records, staff must repeatedly compare information. A payment may be recorded without an updated subscription, or a class list may not reflect its remaining capacity. These inconsistencies motivate a system in which operations share a common database.')
body('Our objective is to provide one application for these activities while retaining clear responsibility for each user. GymVerse separates the public website from the authenticated workspace. Visitors can explore membership plans and upcoming classes, while signed-in users receive pages appropriate to their role. The backend validates requests independently of the interface, and PostgreSQL enforces structural rules that remain valid even when more than one request reaches the system.')
body('The scope includes members, trainers, membership plans, subscriptions, payment records, attendance, exercises, workout plans, workout sessions and class bookings. Dashboards and reports summarize the stored operational data. Optional chat assistance and saved conversations extend the interface. Payment records describe transactions entered into the application; they do not by themselves implement online payment processing or bank settlement.')
table(1,'User roles and responsibilities',['Role','Main responsibility'],[
 ('Administrator','Manage classes, plans, staff approvals, users and reports.'),('Receptionist','Manage member records, subscriptions, payments and attendance.'),('Trainer','Work with assigned members and training plans within permitted access.'),('Member','View permitted personal information, training and class bookings.')],[1.35,5.55])
body('The project applies database concepts to a practical workflow. It demonstrates primary and foreign keys, constrained status values, transactions, indexes, SQL aggregation, audit records and schema migrations. Its success is assessed through functional checks and consistency rules rather than an assumed improvement in gym revenue or member fitness. Those outcomes would require a separate study with real operational data.')

heading('Methodology',newpage=True)
heading('Architecture and technology choices',2)
body('We organized GymVerse as a three-tier application. React renders the browser interface and manages navigation. Axios sends JSON requests to the Express API. The API applies authentication, authorization and validation before issuing parameterized SQL through the PostgreSQL driver. This structure keeps presentation decisions separate from database operations and gives all clients the same server-side rules [1, 2].')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(QA/'architecture.png'),width=Inches(6.8))
caption('Figure 1 GymVerse application architecture')
body('PostgreSQL is the system of record for gym operations. Neon provides the hosted PostgreSQL database, while Render runs the backend. The frontend can be hosted separately or served from the backend after a production build. Optional services have narrower roles: MongoDB stores chat conversations, AI providers generate assistant replies, and SMTP delivers account emails. A failure in optional chat storage is handled with temporary chat behavior.')
body('The development process follows the data model. We define entities and relationships, expose controlled API operations, connect the user interface, and verify the resulting flows. The SQL files remain explicit rather than hidden behind an object-relational mapper. Numbered migrations record applied changes so an existing database can receive upgrades without reloading demonstration data [3].')
body('React Router organizes public and protected pages. Shared interface components provide forms, tables, pagination and feedback states. Lazy loading divides the authenticated application into smaller modules. The visual design uses a dark background, lime accents and gym photography, while membership prices and class availability come from API responses rather than being fixed in the page.')

heading('Database design',2,newpage=True)
body('The schema separates reusable definitions from individual events. A membership plan defines a product; a subscription records a member enrolling in that product for a period. A fitness class defines a session type; a class schedule places it on a date with a trainer and capacity. Class bookings connect members to scheduled sessions [3].')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(QA/'relationships.png'),width=Inches(6.8))
caption('Figure 2 Core membership and class relationships')
table(2,'Core database entities',['Entity group','Stored information'],[
 ('roles and users','Login identities, role links and account state.'),('members and trainers','Profiles and separate links to login accounts.'),('membership_plans and subscriptions','Plan definitions and time-bound member enrolment.'),('payments and attendance','Financial records and member visits.'),('exercises and workout_plans','Reusable exercises and ordered training prescriptions.'),('member_workouts and workout_sessions','Assignments and recorded training sessions.'),('fitness_classes, class_schedules, class_bookings','Class definitions, timetable and reservations.'),('audit_logs and schema_migrations','Change history and applied schema versions.')],[2.75,4.15])
body('Identity primary keys give each record a stable identifier. Foreign keys prevent orphaned references and define deletion behavior. Associative tables handle relationships such as workout plans containing many exercises and members booking many schedules. This separation reduces repeated descriptive data. The diagram shows the central membership and booking relationships rather than every table in the repository.')

heading('Integrity and security',2,newpage=True)
body('The database and API enforce complementary rules. Required fields reject incomplete records, unique values prevent duplicate identities, and check constraints restrict amounts, capacities and status values. The active-subscription partial index permits historical subscriptions while preventing more than one active subscription for a member. API validation adds readable feedback before the database rejects invalid input [3, 4].')
body('Subscription creation uses a transaction. The controller locks the member record, checks the member and selected plan, creates the subscription and opening payment, and commits the related changes together. If an operation fails, rollback prevents a partially completed enrolment. Locking the member row also serializes competing requests for the same member, including when no active subscription row exists yet.')
body('Class booking uses a row lock on the schedule. Within the transaction, the controller checks the session, current bookings and available capacity before reserving a place. Requests for the same class therefore inspect capacity in a controlled order. The locking mechanism is implemented in the controller; a dedicated concurrent load test remains necessary to measure behavior under sustained contention.')
body('Passwords are hashed with bcrypt, and authenticated requests carry a JSON Web Token. Route-level permissions distinguish staff and member operations. Record-level checks then verify ownership or an allowed trainer relationship. This second check matters because a valid member token must not grant access to another member merely by changing an identifier in a URL.')
body('Account recovery uses time-limited tokens, while session invalidation is supported through account token versions and revoked-token records. Request rate limits constrain repeated authentication and chat calls. Parameterized SQL separates data from query syntax. Audit handling removes password hashes before storing change snapshots, and production error responses avoid returning internal stack traces.')
body('These controls address specific risks and do not constitute a complete security certification. Operational protection also requires confidential environment variables, restricted database access, backups and maintained dependencies. The browser currently stores its login token in local storage, so preventing script injection remains especially important. A future authentication review can evaluate HTTP-only cookies and the additional cross-site request protections they require.')

heading('Implementation and deployment',2,newpage=True)
body('The landing page introduces the gym, training programs and facility experience. It reads the plan catalogue and upcoming schedules from public endpoints. Loading, error and empty states keep the interface understandable when data is delayed or no classes have been scheduled. After login, the dashboard and navigation expose role-appropriate actions. The API remains responsible for authorization even when a button is hidden in the browser [2, 4].')
body('The backend groups routes by resource, including authentication, members, trainers, membership plans, subscriptions, payments, attendance, workouts and schedules. Controllers carry out validation-dependent operations, and shared helpers manage transactions, pagination and errors. Dashboards use SQL aggregation so the browser receives summaries rather than downloading all operational records.')
body('The assistant is available to signed-in users and reads the current plan catalogue for membership questions. Provider fallback and time limits handle unavailable AI responses. MongoDB chat history is optional and isolated from the main operational database. The application must not treat assistant output as a replacement for authorized database actions or professional advice.')
body('For the hosted setup, Neon supplies PostgreSQL connection settings and Render runs the Express process. The repository currently reads individual database environment variables. The frontend API address, backend allowed origin and account-link base URL must refer to the deployed domains. A production frontend build is required when its API address changes because Vite embeds that setting during compilation.')
body('Schema upgrades use the migration command, which applies pending numbered SQL files and records their checksums. The required-role migration inserts the four application roles without changing existing role identifiers. Development seed data is separate and must not be replayed over an existing production database. The newer demo timetable provides relative future dates for fresh development installations; administrators create actual classes in the deployed application.')
body('A release check should confirm that the health endpoint reaches PostgreSQL, public plans and schedules load, login succeeds, and protected operations respect user roles. Publishing a backend revision alone does not publish a separately hosted frontend. Production deployment verification and backup restoration exercises should therefore be treated as separate operational checks, not inferred from a successful local build.')

heading('Experiments',newpage=True)
body('Verification is organized into frontend tests, backend regression tests and a full-stack smoke check. Frontend tests exercise application behavior with Vitest. Backend tests use Jest and request-level helpers to check validation and authorization. The smoke script creates an isolated PostgreSQL database, runs migrations and seed data, starts the API, performs checks and removes that test database [5].')
table(3,'Functional verification scenarios',['Scenario','Expected behavior'],[
 ('Correct and incorrect login','Accept valid credentials; reject invalid credentials with a usable message.'),('Member registration','Create a member account and its related profile.'),('Cross-member access','Reject attempts to read or modify another member\'s protected records.'),('Trainer directory privacy','Withhold contact details from members while retaining permitted staff access.'),('Pagination validation','Accept valid pages and reject negative or excessive values.'),('Migrations applied twice','Second execution leaves already applied migrations unchanged.'),('Dashboard revenue','Return numeric values in chronological month order.'),('Audit inspection','Do not retain password hashes in logged snapshots.'),('Development browser origins','Accept the configured localhost origin and its loopback alias.')],[2.05,4.85])
body('The recorded smoke run covered authentication, main read endpoints, ownership regressions, indexes and selected write behavior. Its output provides a repeatable functional baseline. It is not a comprehensive benchmark: there are no measured throughput, latency percentile or production availability results in this report.')
body('Further experiments should run simultaneous reservations for the last class place, subscription races, provider timeouts and expired-session flows. These tests should use disposable databases and controlled accounts. Live Neon records should remain outside any workflow that drops or recreates a database. Email delivery and hosted domain configuration also require environment-specific checks.')

heading('Results',newpage=True)
body('The recorded verification established that the tested local build could initialize its database, start the API and serve the main application workflows. The frontend build also completed. Table 4 reports the observed counts from that verification session; these are recorded results, not a claim that the current hosted deployment was retested [5].')
table(4,'Recorded verification results',['Verification','Observed outcome','Scope'],[
 ('Full-stack smoke check','71 passed, 0 failed','Fresh database, migrations and API checks.'),('Frontend test suite','16 passed','Three frontend test files.'),('Development-origin tests','3 passed','Loopback aliases and production restrictions.'),('Frontend production build','Succeeded','Compiled deployable frontend assets.'),('HTTP origin checks','Both returned HTTP 200','Health requests from localhost and 127.0.0.1 origins.')],[2.15,1.6,3.15])
body('One integration fault was identified during the redesign: the development API accepted localhost but rejected the equivalent 127.0.0.1 browser origin. The origin helper now expands these two loopback names in development while preserving exact production origins. Direct HTTP checks confirmed a successful health response for both addresses.')
body('The timetable also exposed a data issue. Existing demonstration schedules were historical, so the public page correctly had no upcoming entries. Fresh seed data now includes five future sessions using dates relative to installation. Existing databases are not automatically populated with these demonstrations. This distinction prevents a data-loading change from being mistaken for a schema upgrade.')
body('The report supports functional conclusions for the exercised paths. It does not measure usability with gym staff, real payment settlement, long-running production stability or high-volume concurrent use. Optional email, AI and saved chat behavior depend on service credentials and availability. The complete backend suite, including the in-memory MongoDB history tests, is not represented by the three focused origin tests reported here.')

heading('Conclusion and Future Work',newpage=True)
body('GymVerse demonstrates a database-centered approach to gym administration. We connected a role-aware React interface to an Express API and a relational PostgreSQL model. Memberships, payments, attendance, training and bookings share consistent identifiers and relationships. Transactions protect related operations, while database constraints and authorization checks reduce inconsistent updates and unauthorized access.')
body('The principal outcome is an implemented application with a repeatable local verification process. The recorded checks passed for the exercised flows, and the frontend produced a production build. The deployment design uses Neon for PostgreSQL and Render for the backend. Separate configuration and live verification are still required whenever the hosted system is updated.')
heading('Future work',2)
body('The next priority is broader operational testing. Concurrent booking and subscription tests should validate behavior under contention, and monitoring should expose failed requests, database connection failures and service health. Backup restoration should be tested periodically rather than assuming a stored backup is usable.')
body('Payment gateway integration would allow online checkout, but it requires server-side payment verification, idempotent webhook handling and reconciliation with subscription state. The current payment records provide a foundation for that work without implying that external settlement is already automated.')
body('Additional improvements include class reminders, attendance device integrations, richer member progress views and accessibility testing across keyboard, screen reader and mobile use. The schema already contains several supporting record types, but a database field or table alone does not demonstrate an implemented interface or device integration.')
body('The project also offers scope for a focused security review of token storage, session expiry and account recovery. These improvements should be evaluated against clear acceptance tests. Expanding the feature set is useful only when the resulting operations remain understandable to users and consistent in the database.')

p=doc.add_paragraph('References',style='Heading 5'); p.paragraph_format.page_break_before=True
p.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(12); r.bold=True
for text in [
 '[1] GymVerse project repository. PROJECT_GUIDE.md. Architecture, modules, request lifecycle and database design.',
 '[2] GymVerse frontend. src/App.jsx, src/pages/public/Landing.jsx and src/services/api.js. Routing, public data and API integration.',
 '[3] GymVerse database. database/01_schema.sql through database/18_required_roles.sql; gymverse-backend/src/db/migrations.js. Schema, integrity rules, indexes and migrations. Files 07_seed.sql and 08_queries.sql contain development data and example queries rather than schema migrations.',
 '[4] GymVerse backend. src/app.js, src/controllers/subscriptionController.js, src/controllers/scheduleController.js and authentication middleware. API controls, transactions and booking logic.',
 '[5] GymVerse verification. gymverse-backend/verify_stack.js, tests/corsOrigins.test.js and the frontend Vitest suite. Recorded execution results from the project development session.',
 '[6] GymVerse backend and frontend README files and environment examples. Local setup, configuration, scripts and optional service boundaries.'
]: body(text)

temp=QA/'authored.docx'; doc.save(temp)
# Preserve every untouched package part, including the source logo and theme.
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with ZipFile(SOURCE) as src, ZipFile(temp) as built, ZipFile(OUT,'w',ZIP_DEFLATED) as out:
    for n in src.namelist(): out.writestr(n,built.read(n) if n in editable else src.read(n))
    for n in built.namelist():
        if n not in src.namelist(): out.writestr(n,built.read(n))
with ZipFile(OUT) as z:
    changed=[n for n,h in inventory.items() if hashlib.sha256(z.read(n)).hexdigest()!=h]
    assert set(changed)<=editable,changed
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
print(OUT)
print('Source preserved; changed existing package parts:',changed)
print('Sections:',len(Document(OUT).sections))
